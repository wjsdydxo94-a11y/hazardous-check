from flask import Flask, render_template, request, redirect, url_for, send_file
import pandas as pd
from datetime import datetime, timedelta
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
DATA_FILE = 'inspection_db.csv'

def init_db():
    if not os.path.exists(DATA_FILE):
        df = pd.DataFrame(columns=[
            "점검일시", "업체명", "점검자", "건축물구조", "소방환기", 
            "표지저장", "온습도누출", "온도", "습도", "특이사항"
        ])
        df.to_csv(DATA_FILE, index=False, encoding='utf-8-sig')

def generate_monthly_excel(company_name):
    init_db()
    df_db = pd.read_csv(DATA_FILE)
    
    if '업체명' in df_db.columns:
        df_db = df_db[df_db['업체명'] == company_name]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "위험물일일점검표"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="맑은 고딕", size=16, bold=True)
    font_header = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="맑은 고딕", size=10)

    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_meta = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # 1. Title
    ws.merge_cells("A1:G1")
    ws["A1"] = "위험물 일일점검표"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 40

    # 2. Sub-note
    ws.merge_cells("A2:G2")
    ws["A2"] = "* 본 점검표는 일진소방(주) 양식 및 위험물안전관리법 기준에 의거하여 매일 현장 점검 후 기록관리하는 법정 서식입니다."
    ws["A2"].font = Font(name="맑은 고딕", size=9, italic=True, color="595959")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    # 3. Metadata Header Box (안전관리자란 공백 처리하여 수기 작성 가능)
    current_month_str = datetime.now().strftime('%Y년 %m월')
    metadata = [
        ("사업장명", f"주식회사 {company_name}", "점검년월", current_month_str),
        ("점검대상", "옥내저장소", "안전관리자", "")
    ]

    for r_idx, meta in enumerate(metadata, 3):
        ws.row_dimensions[r_idx].height = 24
        ws.merge_cells(start_row=r_idx, start_column=1, end_row=r_idx, end_column=2)
        ws.cell(row=r_idx, column=1, value=meta[0]).fill = fill_meta
        ws.cell(row=r_idx, column=1).font = font_body
        ws.cell(row=r_idx, column=1).alignment = align_center
        ws.cell(row=r_idx, column=1).border = thin_border
        ws.cell(row=r_idx, column=2).border = thin_border

        ws.merge_cells(start_row=r_idx, start_column=3, end_row=r_idx, end_column=4)
        ws.cell(row=r_idx, column=3, value=meta[1]).alignment = align_center
        ws.cell(row=r_idx, column=3).font = font_body
        ws.cell(row=r_idx, column=3).border = thin_border
        ws.cell(row=r_idx, column=4).border = thin_border

        ws.cell(row=r_idx, column=5, value=meta[2]).fill = fill_meta
        ws.cell(row=r_idx, column=5).font = font_body
        ws.cell(row=r_idx, column=5).alignment = align_center
        ws.cell(row=r_idx, column=5).border = thin_border

        ws.merge_cells(start_row=r_idx, start_column=6, end_row=r_idx, end_column=7)
        ws.cell(row=r_idx, column=6, value=meta[3]).alignment = align_center
        ws.cell(row=r_idx, column=6).font = font_body
        ws.cell(row=r_idx, column=6).border = thin_border
        ws.cell(row=r_idx, column=7).border = thin_border

    ws.row_dimensions[5].height = 10

    # 4. Table Headers (일진소방 양식 세부 항목 반영)
    headers = [
        "일자", 
        "1. 건축물구조/피뢰\n(균열/피뢰설비)", 
        "2. 소방/환기배출\n(소화기/환기팬)", 
        "3. 표지판/저장취급\n(금연표지/관리상태)", 
        "4. 온습도/누출방지\n(온습도/누출비산방지)", 
        "점검 결과", 
        "점검자 서명\n(선임자)"
    ]

    ws.row_dimensions[6].height = 30
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=6, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border

    # 5. Populate Days (1 to 31)
    for day in range(1, 32):
        r_idx = 6 + day
        ws.row_dimensions[r_idx].height = 22
        
        date_prefix = f"{datetime.now().strftime('%m월')} {day:02d}일"
        
        matched_row = None
        for idx, row in df_db.iterrows():
            try:
                dt = datetime.strptime(str(row['점검일시']), '%Y-%m-%d %H:%M:%S')
                current_now = datetime.now()
                if dt.year == current_now.year and dt.month == current_now.month and dt.day == day:
                    matched_row = row
                    break
            except Exception:
                if f"-{day:02d} " in str(row['점검일시']) or str(row['점검일시']).endswith(f"-{day:02d}"):
                    matched_row = row
                    break

        if matched_row is not None:
            def format_status(val):
                val_str = str(val).strip()
                if any(kw in val_str for kw in ['양호', '정상', '이상없음', '적합']):
                    return "O"
                else:
                    return "X"

            v1 = format_status(matched_row.get('건축물구조', '양호'))
            v2 = format_status(matched_row.get('소방환기', '양호'))
            v3 = format_status(matched_row.get('표지저장', '양호'))
            v4 = format_status(matched_row.get('온습도누출', '양호'))
            
            if v1 == "O" and v2 == "O" and v3 == "O" and v4 == "O":
                v_result = "적합 (양호)"
            else:
                v_result = "부적합 (정비요함)"
                
            insp = str(matched_row.get('점검자', ''))
            v_signer = f"{company_name} / {insp}" if insp else company_name
        else:
            v1 = "[   ]"
            v2 = "[   ]"
            v3 = "[   ]"
            v4 = "[   ]"
            v_result = "[   ]"
            v_signer = ""

        row_vals = [date_prefix, v1, v2, v3, v4, v_result, v_signer]
        
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = font_body
            cell.alignment = align_center
            cell.border = thin_border
            if day % 2 == 1:
                cell.fill = PatternFill(start_color="FAFAFA", end_color="FAFAFA", fill_type="solid")

    # 6. Remarks Section
    r_remark = 39
    ws.row_dimensions[r_remark].height = 25
    ws.merge_cells(start_row=r_remark, start_column=1, end_row=r_remark, end_column=7)
    ws.cell(row=r_remark, column=1, value="※ 점검지시 및 특이사항 기록").font = Font(name="맑은 고딕", size=10, bold=True)
    ws.cell(row=r_remark, column=1).alignment = Alignment(horizontal="left", vertical="center")

    for r in range(40, 44):
        ws.row_dimensions[r].height = 24
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = thin_border

    col_widths = [14, 22, 22, 22, 20, 16, 22]
    for idx, width in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    report_filename = f'위험물일일점검표_{company_name}.xlsx'
    wb.save(report_filename)
    return report_filename

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/submit', methods=['POST'])
def submit():
    init_db()
    company = request.form.get('company', '아로마솔루션')
    inspector = request.form.get('inspector')
    item1 = request.form.get('item1', '양호')
    item2 = request.form.get('item2', '양호')
    item3 = request.form.get('item3', '양호')
    item4 = request.form.get('item4', '양호')
    temp = request.form.get('temp')
    humidity = request.form.get('humidity')
    remarks = request.form.get('remarks', '-')
    
    now = (datetime.utcnow() + timedelta(hours=9)).strftime('%Y-%m-%d %H:%M:%S')
    
    new_row = pd.DataFrame([{
        "점검일시": now,
        "업체명": company,
        "점검자": inspector,
        "건축물구조": item1,
        "소방환기": item2,
        "표지저장": item3,
        "온습도누출": item4,
        "온도": temp,
        "습도": humidity,
        "특이사항": remarks
    }])
    
    df = pd.read_csv(DATA_FILE)
    df = pd.concat([df, new_row], ignore_index=True)
    df.to_csv(DATA_FILE, index=False, encoding='utf-8-sig')
    
    return render_template('success.html')

@app.route('/admin')
def admin():
    init_db()
    df = pd.read_csv(DATA_FILE)
    records = df.to_dict(orient='records')
    return render_template('admin.html', records=records)

@app.route('/download')
def download():
    company = request.args.get('company', '아로마솔루션')
    report_filename = generate_monthly_excel(company)
    return send_file(report_filename, as_attachment=True)

if __name__ == '__main__':
    init_db()
    port = int(os.environ.get("PORT", 5000))
    app.run(host='0.0.0.0', port=port)
